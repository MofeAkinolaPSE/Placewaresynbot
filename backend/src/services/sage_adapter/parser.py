from __future__ import annotations

import csv
import io
import json
from typing import List, Dict, Any


def parse_csv(content: bytes) -> List[Dict[str, Any]]:
    text = content.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    return [dict(row) for row in reader]


def parse_json(content: bytes) -> List[Dict[str, Any]]:
    raw = json.loads(content.decode("utf-8", errors="replace"))
    if isinstance(raw, dict):
        if isinstance(raw.get("rows"), list):
            raw = raw["rows"]
        else:
            raw = [raw]
    if not isinstance(raw, list):
        raise ValueError("JSON payload must be an object or list of objects")
    rows: List[Dict[str, Any]] = []
    for item in raw:
        if isinstance(item, dict):
            rows.append(dict(item))
    return rows


def parse_xlsx(content: bytes) -> List[Dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except Exception as exc:
        raise ValueError("XLSX support unavailable. Install openpyxl.") from exc

    workbook = load_workbook(filename=io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    rows_iter = sheet.iter_rows(values_only=True)
    try:
        first = next(rows_iter)
    except StopIteration:
        return []
    headers = [str(col).strip() if col is not None else "" for col in first]
    parsed: List[Dict[str, Any]] = []
    for row in rows_iter:
        values = list(row)
        if not any(v is not None and str(v).strip() != "" for v in values):
            continue
        parsed.append({headers[i]: values[i] if i < len(values) else None for i in range(len(headers))})
    return parsed


def parse_tabular_upload(content: bytes, filename: str | None = None, content_type: str | None = None) -> List[Dict[str, Any]]:
    filename = (filename or "").lower()
    content_type = (content_type or "").lower()

    if filename.endswith(".csv") or "text/csv" in content_type:
        return parse_csv(content)
    if filename.endswith(".json") or "json" in content_type:
        return parse_json(content)
    if filename.endswith(".xlsx") or "spreadsheetml" in content_type:
        return parse_xlsx(content)

    stripped = content.lstrip()
    if stripped.startswith((b"{", b"[")):
        return parse_json(content)
    if b"," in content[:512]:
        return parse_csv(content)

    raise ValueError("Unsupported file format. Allowed formats: .csv, .json, .xlsx")
