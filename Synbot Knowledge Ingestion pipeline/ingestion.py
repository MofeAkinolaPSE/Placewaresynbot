"""SynBot Knowledge Ingestion Pipeline.

Ingests DOCX, PDF, TXT, CSV, and MD files into the knowledge_chunks and
ingested_documents tables in PostgreSQL (pgvector).

Embeddings are produced locally via FastEmbed (all-MiniLM-L6-v2, 384-dim) —
identical model to backend/src/embed_proxy.py so ingest and query vectors are
compatible.

Usage (CLI):
    python ingestion.py --folder "placeware data/"
    python ingestion.py --file path/to/document.pdf
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import logging
import os
from pathlib import Path

import psycopg2
from dotenv import load_dotenv
from pypdf import PdfReader
from docx import Document

# ================================
# CONFIG
# ================================

load_dotenv()

# Reads DATABASE_URL from .env — same variable the backend uses.
# Falls back to individual POSTGRES_* vars for local dev convenience.
DATABASE_URL = os.getenv("DATABASE_URL") or (
    "postgresql://{user}:{password}@{host}:{port}/{dbname}".format(
        user=os.getenv("POSTGRES_USER", "postgres"),
        password=os.getenv("POSTGRES_PASSWORD", "password"),
        host=os.getenv("POSTGRES_HOST", "localhost"),
        port=os.getenv("POSTGRES_PORT", "5432"),
        dbname=os.getenv("POSTGRES_DB", "placeware"),
    )
)

CHUNK_SIZE    = 800
CHUNK_OVERLAP = 150

logging.basicConfig(level=logging.INFO, format="%(levelname)s  %(message)s")
logger = logging.getLogger(__name__)

# ================================
# EMBEDDING (FastEmbed — local, no API key)
# ================================

try:
    from fastembed import TextEmbedding as _FastEmbed
    _embed_model = _FastEmbed(model_name="sentence-transformers/all-MiniLM-L6-v2")
    logger.info("FastEmbed model loaded.")
except Exception as _e:
    _embed_model = None
    logger.error("FastEmbed init failed: %s", _e)


def create_embedding(text: str) -> list[float] | None:
    """Return a 384-dim embedding vector using the local FastEmbed model."""
    if _embed_model is None:
        logger.error("Embedding model not available.")
        return None
    try:
        return list(_embed_model.embed([text]))[0].tolist()
    except Exception as exc:
        logger.error("Embedding failed: %s", exc)
        return None


# ================================
# DATABASE CONNECTION
# ================================

def get_db():
    """Open a psycopg2 connection using DATABASE_URL."""
    return psycopg2.connect(DATABASE_URL)

# ================================
# DOCUMENT HASHING
# ================================

def generate_file_hash(filepath: str) -> str:
    """SHA-256 of raw file bytes — used for deduplication."""
    sha256 = hashlib.sha256()
    with open(filepath, "rb") as f:
        while chunk := f.read(8192):
            sha256.update(chunk)
    return sha256.hexdigest()


def hash_text(text: str) -> str:
    """SHA-256 of a text string — used when ingesting generated documents."""
    return hashlib.sha256(text.encode("utf-8")).hexdigest()

# ================================
# DOCUMENT LOADERS
# ================================

def load_pdf(path: str) -> str:
    """Extract text from PDF page-by-page, logging pages that fail."""
    reader = PdfReader(path)
    parts: list[str] = []
    for i, page in enumerate(reader.pages):
        try:
            text = page.extract_text() or ""
            if text.strip():
                parts.append(text.strip())
        except Exception as exc:
            logger.warning("PDF page %d skipped (%s): %s", i + 1, Path(path).name, exc)
    content = "\n".join(parts)
    if not content.strip():
        logger.warning(
            "%s has no extractable text — likely a scanned/image PDF. "
            "Add pytesseract for OCR support.",
            Path(path).name,
        )
    return content


def load_docx(path: str) -> str:
    """Extract ALL text from a DOCX file.

    Three-pass strategy:
      Pass 1 — XML w:p traversal of body (nested tables + text boxes)
      Pass 2 — Headers and footers for every section
      Pass 3 — Content Controls (w:sdt) at body level
    Deduplicates lines so merged table cells don't repeat.
    Falls back to doc.paragraphs API if all passes yield nothing.
    """
    from docx.oxml.ns import qn
    doc  = Document(path)
    seen: set[str] = set()
    parts: list[str] = []

    def _collect(root) -> None:
        for para in root.iter(qn("w:p")):
            text = "".join(node.text for node in para.iter(qn("w:t")) if node.text).strip()
            if text and text not in seen:
                seen.add(text)
                parts.append(text)

    # Pass 1: full body tree
    _collect(doc.element.body)

    # Pass 2: headers and footers
    for section in doc.sections:
        for hf in (
            section.header, section.footer,
            section.even_page_header, section.even_page_footer,
            section.first_page_header, section.first_page_footer,
        ):
            try:
                if hf and hf._element is not None:
                    _collect(hf._element)
            except Exception:
                pass

    # Pass 3: Content Controls (w:sdt)
    for sdt in doc.element.body.iter(qn("w:sdt")):
        text = "".join(node.text for node in sdt.iter(qn("w:t")) if node.text).strip()
        if text and text not in seen:
            seen.add(text)
            parts.append(text)

    # Fallback
    if not parts:
        for para in doc.paragraphs:
            t = para.text.strip()
            if t and t not in seen:
                seen.add(t)
                parts.append(t)

    result = "\n".join(parts)
    if not result.strip():
        logger.warning(
            "%s yielded no text — may be encrypted or use unsupported formatting "
            "(SmartArt, embedded objects). Re-save as plain .docx to fix.",
            Path(path).name,
        )
    return result


def load_txt(path: str) -> str:
    with open(path, "r", encoding="utf-8") as f:
        return f.read()


def load_csv(path: str) -> str:
    """Convert CSV rows to readable key=value text blocks."""
    lines: list[str] = []
    with open(path, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            lines.append("  |  ".join(f"{k}: {v}" for k, v in row.items()))
    return "\n".join(lines)


def load_xlsx(path: str) -> str:
    """Extract all cell text from every sheet in an Excel workbook."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in wb.worksheets:
        parts.append(f"Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                parts.append("  |  ".join(cells))
    wb.close()
    return "\n".join(parts)


def load_md(path: str) -> str:
    """Read markdown as plain text — headings preserved as natural structure."""
    with open(path, "r", encoding="utf-8") as f:
        return f.read()


def load_document(filepath: str) -> str:
    ext = Path(filepath).suffix.lower()
    if ext == ".pdf":
        return load_pdf(filepath)
    if ext == ".docx":
        return load_docx(filepath)
    if ext in (".txt",):
        return load_txt(filepath)
    if ext == ".csv":
        return load_csv(filepath)
    if ext in (".md", ".markdown"):
        return load_md(filepath)
    if ext in (".xlsx", ".xls"):
        return load_xlsx(filepath)
    raise ValueError(f"Unsupported file format: {ext}")

# ================================
# SMART CHUNKING
# ================================

def chunk_text(text: str) -> list[str]:
    """Fixed-size chunks with overlap.  Skips whitespace-only chunks."""
    chunks: list[str] = []
    start = 0
    while start < len(text):
        chunk = text[start : start + CHUNK_SIZE].strip()
        if chunk:
            chunks.append(chunk)
        start += CHUNK_SIZE - CHUNK_OVERLAP
    return chunks

# ================================
# METADATA DETECTION
# ================================

_TYPE_KEYWORDS: list[tuple[list[str], str, str]] = [
    (["audit"],                                    "audit_report",       "quality"),
    (["deviation", "capa"],                        "deviation_report",   "quality"),
    (["maintenance", "calibr", "schedule"],        "maintenance_log",    "operations"),
    (["fumigation", "pest", "sanitiz", "hygiene"], "environmental_log",  "operations"),
    (["sop", "procedure", "index"],                "sop",                "quality"),
    (["recall"],                                   "recall_notice",      "quality"),
    (["invoice", "sales"],                         "sales_document",     "finance"),
    (["inventory", "stock"],                       "inventory_report",   "operations"),
    (["staff", "payroll", "hr"],                   "hr_document",        "hr"),
]


def detect_metadata(file_name: str, source: str = "manual") -> dict:
    """Infer document_type and department from the filename."""
    name_lower = file_name.lower()
    for keywords, doc_type, dept in _TYPE_KEYWORDS:
        if any(kw in name_lower for kw in keywords):
            return {"document_type": doc_type, "department": dept, "source": source}
    return {"document_type": "general_document", "department": "general", "source": source}

# ================================
# DUPLICATE CHECK
# ================================

def document_exists(file_hash: str) -> bool:
    conn = get_db()
    cur  = conn.cursor()
    cur.execute(
        "SELECT 1 FROM ingested_documents WHERE file_hash = %s LIMIT 1",
        (file_hash,),
    )
    exists = cur.fetchone() is not None
    cur.close()
    conn.close()
    return exists

# ================================
# SAVE DOCUMENT RECORD
# ================================

def register_document(
    document_id: str,
    file_name: str,
    file_path: str,
    file_hash: str,
    metadata: dict,
) -> None:
    conn = get_db()
    cur  = conn.cursor()
    cur.execute(
        """
        INSERT INTO ingested_documents
            (document_id, file_name, file_path, file_hash,
             document_type, department, source, metadata)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
        ON CONFLICT (file_hash) DO NOTHING
        """,
        (
            document_id,
            file_name,
            file_path,
            file_hash,
            metadata.get("document_type", "general_document"),
            metadata.get("department", "general"),
            metadata.get("source", "manual"),
            json.dumps(metadata),
        ),
    )
    conn.commit()
    cur.close()
    conn.close()

# ================================
# SAVE CHUNKS
# ================================

def save_chunks_batch(
    document_id: str,
    chunks: list[str],
    embeddings: list[list[float] | None],
    metadata: dict,
) -> int:
    """Insert all chunks for one document in a single DB transaction.

    Much faster than one connection-per-chunk and guarantees atomicity
    (either all chunks land or none do).
    """
    conn = get_db()
    cur  = conn.cursor()
    try:
        for i, (chunk, embedding) in enumerate(zip(chunks, embeddings)):
            cur.execute(
                """
                INSERT INTO knowledge_chunks
                    (document_id, chunk_index, content, embedding, metadata)
                VALUES (%s, %s, %s, %s, %s)
                """,
                (
                    document_id,
                    i,
                    chunk,
                    embedding,
                    json.dumps(metadata),
                ),
            )
        conn.commit()
        return len(chunks)
    except Exception:
        conn.rollback()
        raise
    finally:
        cur.close()
        conn.close()

# ================================
# INGESTION PIPELINE
# ================================

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".csv", ".md", ".markdown", ".xlsx", ".xls"}


def ingest_document(filepath: str, source: str = "manual") -> dict:
    """Ingest a single file.  Returns a summary dict for programmatic callers."""
    filepath = str(Path(filepath).resolve())
    logger.info("Processing %s", filepath)

    file_hash = generate_file_hash(filepath)

    if document_exists(file_hash):
        logger.info("Already ingested (hash match): %s", Path(filepath).name)
        return {"status": "skipped", "reason": "already_ingested", "file": filepath}

    file_name   = Path(filepath).name
    document_id = file_hash[:16]
    metadata    = detect_metadata(file_name, source=source)

    try:
        text = load_document(filepath)
    except Exception as exc:
        logger.error("Failed to load %s: %s", filepath, exc)
        return {"status": "error", "reason": str(exc), "file": filepath}

    logger.info("Extracted %d chars from %s", len(text), file_name)

    chunks = chunk_text(text)
    if not chunks:
        logger.warning("No content extracted from %s", filepath)
        return {"status": "skipped", "reason": "empty_content", "file": filepath}

    logger.info("Chunked into %d chunk(s) — embedding now…", len(chunks))

    register_document(
        document_id=document_id,
        file_name=file_name,
        file_path=filepath,
        file_hash=file_hash,
        metadata=metadata,
    )

    # Compute all embeddings up-front, then write in one transaction
    embeddings: list[list[float] | None] = [
        create_embedding(chunk) for chunk in chunks
    ]

    try:
        ingested = save_chunks_batch(document_id, chunks, embeddings, metadata)
    except Exception as exc:
        logger.error("Failed to save chunks for %s: %s", file_name, exc)
        return {"status": "error", "reason": str(exc), "file": filepath}

    logger.info("Ingested %d chunks from %s", ingested, file_name)
    return {"status": "ok", "chunks": ingested, "document_id": document_id, "file": filepath}


# ================================
# RUNNER
# ================================

def ingest_folder(folder: str, source: str = "manual") -> list[dict]:
    """Ingest all supported files in a directory (non-recursive)."""
    results: list[dict] = []
    folder_path = Path(folder)

    if not folder_path.is_dir():
        logger.error("Not a directory: %s", folder)
        return results

    files = [
        p for p in folder_path.iterdir()
        if p.is_file() and p.suffix.lower() in SUPPORTED_EXTENSIONS
    ]

    if not files:
        logger.warning("No supported files found in %s", folder)
        return results

    logger.info("Found %d file(s) to process in %s", len(files), folder)
    for file_path in sorted(files):
        try:
            result = ingest_document(str(file_path), source=source)
        except Exception as exc:
            logger.error("Unexpected error processing %s: %s", file_path.name, exc)
            result = {"status": "error", "reason": str(exc), "file": str(file_path)}
        results.append(result)

    ok       = sum(1 for r in results if r["status"] == "ok")
    skipped  = sum(1 for r in results if r["status"] == "skipped")
    errors   = sum(1 for r in results if r["status"] == "error")
    logger.info("Done — ok: %d, skipped: %d, errors: %d", ok, skipped, errors)
    return results


# ================================
# MAIN
# ================================

if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="SynBot Knowledge Ingestion CLI"
    )
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument(
        "--folder", metavar="DIR",
        help="Ingest all supported files in this directory",
    )
    group.add_argument(
        "--file", metavar="FILE",
        help="Ingest a single file",
    )
    parser.add_argument(
        "--source",
        default="manual",
        help="Source tag stored in metadata (default: manual)",
    )
    args = parser.parse_args()

    if args.file:
        result = ingest_document(args.file, source=args.source)
        print(result)
    else:
        results = ingest_folder(args.folder, source=args.source)
        for r in results:
            print(r)